import openpyxl
wb_test = openpyxl.Workbook()
ws_test = wb_test.active
ws_test.append(["Öğrenci Adı", "Vize", "Final"])
ws_test.append(["Ahmet", 60, 80])
ws_test.append(["Mehmet", 40, 50])
ws_test.append(["Ayşe", 80, 90])
wb_test.save("ogrenci_notlari.xlsx")

dosya = openpyxl.load_workbook("ogrenci_notlari.xlsx")
sayfa = dosya.active

sayfa.cell(row=1, column=4, value="Ortalama")

satir_sayisi = sayfa.max_row

for satir in range(2, satir_sayisi + 1):
    vize = sayfa.cell(row=satir, column=2).value
    final = sayfa.cell(row=satir, column=3).value
    
    # Ortalama hesaplama (%40 vize + %60 final)
    ortalama = (vize * 0.4) + (final * 0.6)
    
    sayfa.cell(row=satir, column=4, value=ortalama)

dosya.save("ogrenci_notlari_guncel.xlsx")
print("1. Proje Başarıyla Çalıştı!")
