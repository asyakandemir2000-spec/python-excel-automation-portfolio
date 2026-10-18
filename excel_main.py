import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Veri Kaynağının Oluşturulması
data = {
    "Product ID": [1, 2, 3, 4, 5],
    "Product Name": [
        "Wireless Mechanical Keyboard",
        "Ergonomic USB-C Gaming Mouse",
        "UltraWide 4K Monitor 32 Inch",
        "Thunderbolt 4 Docking Station",
        "Noise Cancelling Headphones"
    ],
    "Inventory Status": [
        "In Stock",
        "Out of Stock",
        "In Stock",
        "Out of Stock",
        "In Stock"
    ]
}

df = pd.DataFrame(data)

# 2. Excel Çalışma Kitabı Oluşturma
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Inventory Status"
ws.views.sheetView[0].showGridLines = True

# 3. Başlık Stilleri (Mavi Dolgu, Beyaz Yazı)
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Başlıkları Ekleme
headers = ["Product ID", "Product Name", "Inventory Status"]
ws.append(headers)

for col_num in range(1, 4):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center" if col_num != 2 else "left", vertical="center")
    cell.border = thin_border

# 4. Verileri Yazma ve Koşullu Renklendirme
in_stock_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
in_stock_font = Font(name="Calibri", size=11, color="375623", bold=True)
out_stock_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
out_stock_font = Font(name="Calibri", size=11, color="C65911", bold=True)

for row_idx, row in df.iterrows():
    row_num = row_idx + 2
    ws.append([row["Product ID"], row["Product Name"], row["Inventory Status"]])
    
    c1 = ws.cell(row=row_num, column=1)
    c2 = ws.cell(row=row_num, column=2)
    c3 = ws.cell(row=row_num, column=3)
    
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c2.alignment = Alignment(horizontal="left", vertical="center")
    c3.alignment = Alignment(horizontal="center", vertical="center")
    
    for c in (c1, c2, c3):
        c.border = thin_border
            
    if row["Inventory Status"] == "In Stock":
        c3.fill = in_stock_fill
        c3.font = in_stock_font
    else:
        c3.fill = out_stock_fill
        c3.font = out_stock_font

# 5. Özet Bilgileri Ekleme
ws.append(["", "", ""]) # Boş satır

summary_rows = [
    ("Total Products", len(df)),
    ("In Stock", len(df[df["Inventory Status"] == "In Stock"])),
    ("Out of Stock", len(df[df["Inventory Status"] == "Out of Stock"]))
]

summary_font = Font(name="Calibri", size=11, bold=True)

for label, val in summary_rows:
    current_row = ws.max_row + 1
    ws.cell(row=current_row, column=1, value=label).font = summary_font
    ws.cell(row=current_row, column=2, value=val).font = summary_font
    ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=current_row, column=2).alignment = Alignment(horizontal="center", vertical="center")

# 6. Sütun Genişliklerini Ayarlama
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 15)

# 7. Excel Dosyasını Kaydetme
file_name = "inventory_report.xlsx"
wb.save(file_name)
print(f"Başarılı! Dosya oluşturuldu: {file_name}")